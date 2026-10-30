"""Pruebas del modelo de Excel para importar productos (Admin > Productos > Importar Excel > Descargar modelo) y de lo
que el importador tolera: encabezados con tildes, el libro guardado en otra hoja, filas en blanco y codigos como numero.

Se ejecutan contra una base SQLite temporal, nunca contra la base real:

    python -m unittest test_modelo_excel -v
"""
import io
import os
import re
import tempfile
import unittest
import uuid

_tmpdir = tempfile.mkdtemp(prefix='dewill_test_')
os.environ['DATABASE_URL'] = 'sqlite:///' + os.path.join(_tmpdir, 'test.db').replace(os.sep, '/')

from openpyxl import Workbook, load_workbook  # noqa: E402  (la URL de la base debe fijarse antes de importar la app)
from app import app, db  # noqa: E402
from import_template import IMPORT_COLUMNS  # noqa: E402
from models import Brand, Presentation, Product, ProductVariantCode  # noqa: E402

app.config['TESTING'] = True
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class ModeloExcelTest(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()
        self.client.post('/admin/login', data={'username': 'admin', 'password': 'admin123'})
        self.s = uuid.uuid4().hex[:6].upper()
        with app.app_context():
            db.session.add(Brand(name=f'MARCA {self.s}', slug=f'marca-{self.s.lower()}'))
            db.session.add(Presentation(name=f'{self.s} GL', slug=f'{self.s.lower()}-gl'))
            db.session.commit()

    def modelo(self):
        resp = self.client.get('/admin/productos/importar/modelo')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.mimetype, XLSX)
        self.assertIn('modelo_importacion_productos.xlsx', resp.headers['Content-Disposition'])
        return load_workbook(io.BytesIO(resp.data))

    def importar(self, wb):
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = self.client.post('/admin/productos/importar', data={'excel_file': (buf, 'productos.xlsx')},
                                content_type='multipart/form-data', follow_redirects=True)
        self.assertEqual(resp.status_code, 200)
        return re.sub(r'\s+', ' ', resp.get_data(as_text=True))

    def fila(self, codigo, producto, color=None, presentacion=None):
        return [codigo, 'LINEA DECORATIVA', 'PINTURA LATEX', f'MARCA {self.s}', producto, color, presentacion]

    def codigos(self, nombre):
        with app.app_context():
            p = Product.query.filter_by(name=nombre).one()
            pres = {x.id: x.name for x in Presentation.query.all()}
            return p.sku, {(pres.get(c.presentation_id), c.color_name): c.code for c in p.variant_codes}

    # ---------- el modelo ----------

    def test_el_modelo_trae_la_hoja_que_se_importa_instrucciones_y_listas(self):
        wb = self.modelo()
        self.assertEqual(wb.sheetnames, ['Productos', 'Instrucciones', 'Listas'])
        self.assertEqual(wb.active.title, 'Productos')
        hoja = wb['Productos']
        self.assertEqual([c.value for c in hoja[1]], IMPORT_COLUMNS)
        self.assertEqual(hoja['A2'].number_format, '@', 'CODIGO como texto para conservar los ceros')
        self.assertGreaterEqual(len(hoja.data_validations.dataValidation), 3)
        listas = {c.value for col in wb['Listas'].iter_cols(min_row=2) for c in col if c.value}
        self.assertIn(f'MARCA {self.s}', listas)
        self.assertIn(f'{self.s} GL', listas)

    def test_el_ejemplo_usa_codigos_reales_del_catalogo(self):
        self.importar(self.llenar(self.modelo(), [self.fila(f'E{self.s}', f'EJEMPLO {self.s}', 'AZUL #7', f'{self.s} GL')]))
        hoja = self.modelo()['Instrucciones']
        filas = [[c.value for c in row] for row in hoja.iter_rows()]
        inicio = next(i for i, f in enumerate(filas) if f[1:8] == IMPORT_COLUMNS)   # encabezado del ejemplo
        ejemplo = [f[1:8] for f in filas[inicio + 1:] if any(f[1:8])]
        self.assertTrue(ejemplo)
        with app.app_context():
            reales = {c.code for c in ProductVariantCode.query.all()}
        self.assertTrue(all(f[0] in reales for f in ejemplo), ejemplo)
        con_codigo_de_color = [f for f in ejemplo if f[5] and '#' in f[5]]
        self.assertTrue(con_codigo_de_color, 'el ejemplo muestra un color con su codigo de fabrica')

    # ---------- de ida y vuelta ----------

    def llenar(self, wb, filas, hoja='Productos'):
        ws = wb[hoja]
        for r, fila in enumerate(filas, start=2):
            for c, valor in enumerate(fila, start=1):
                ws.cell(row=r, column=c, value=valor)
        return wb

    def test_el_modelo_lleno_se_importa_y_conserva_los_ceros(self):
        producto = f'LATEX MODELO {self.s}'
        mensaje = self.importar(self.llenar(self.modelo(), [
            self.fila('0025', producto, 'ALMENDRA', f'{self.s} GL'),
            self.fila('0026', producto, 'AMARILLO OCRE #12', f'{self.s} GL'),
        ]))
        self.assertIn('2 productos importados, 0 omitidos', mensaje, 'las filas preparadas en blanco no cuentan')
        sku, codigos = self.codigos(producto)
        self.assertEqual(sku, '0025')
        self.assertEqual(codigos, {(f'{self.s} GL', 'ALMENDRA'): '0025', (f'{self.s} GL', 'AMARILLO OCRE'): '0026'})

    def test_se_importa_aunque_se_guarde_mirando_otra_hoja(self):
        producto = f'LATEX OTRA HOJA {self.s}'
        wb = self.llenar(self.modelo(), [self.fila('7001', producto, 'BLANCO', f'{self.s} GL')])
        wb.active = 1   # el libro se guardo con las instrucciones abiertas
        self.importar(wb)
        self.assertEqual(self.codigos(producto)[1], {(f'{self.s} GL', 'BLANCO'): '7001'})

    def test_encabezados_con_tildes_y_codigo_escrito_como_numero(self):
        producto = f'LATEX TILDES {self.s}'
        wb = Workbook()
        ws = wb.active
        ws.append(['CÓDIGO', 'CATEGORÍA', 'SUB CATEGORÍA', 'MARCA', 'PRODUCTO', 'COLOR', 'PRESENTACIÓN'])
        ws.append([5002.0, 'LINEA DECORATIVA', 'PINTURA LATEX', f'MARCA {self.s}', producto, 'BLANCO', f'{self.s} GL'])
        self.importar(wb)
        self.assertEqual(self.codigos(producto), ('5002', {(f'{self.s} GL', 'BLANCO'): '5002'}))

    # ---------- acceso ----------

    def test_sin_sesion_no_se_descarga(self):
        resp = app.test_client().get('/admin/productos/importar/modelo')
        self.assertEqual(resp.status_code, 302)

    def test_la_pagina_de_importar_y_la_lista_ofrecen_el_modelo(self):
        for path in ('/admin/productos/importar', '/admin/productos'):
            self.assertIn('/admin/productos/importar/modelo', self.client.get(path).get_data(as_text=True))


if __name__ == '__main__':
    unittest.main()
