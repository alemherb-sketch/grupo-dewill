"""Modelo de Excel para Admin > Productos > Importar Excel.

El libro tiene la hoja que lee el importador (Productos), las instrucciones con un ejemplo y las listas con los nombres
que ya usa la web: el importador reconoce cada producto por su nombre y su marca, y cada presentacion por su nombre,
asi que escribirlos igual que en la web evita productos duplicados o presentaciones nuevas por error.
"""
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# Columnas tal como las lee el importador (mayusculas o minusculas y tildes dan igual: ver import_header)
IMPORT_COLUMNS = ['CODIGO', 'CATEGORIA', 'SUB CATEGORIA', 'MARCA', 'PRODUCTO', 'COLOR', 'PRESENTACION']
IMPORT_REQUIRED = ['Categoria', 'Sub Categoria', 'Marca', 'Producto']
TEMPLATE_ROWS = 1000   # filas preparadas con CODIGO como texto y listas desplegables

HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='0F2B5B')
WIDTHS = [12, 26, 30, 22, 46, 30, 18]

# Ejemplo si el catalogo aun no tiene codigos (los nombres son los del Excel de la web)
DEFAULT_EXAMPLES = [
    ['0025', 'LINEA DECORATIVA', 'PINTURA LATEX', 'AMERICAN COLORS', 'LATEX AMERICAN COLORS', 'ALMENDRA', '1GL'],
    ['5002', 'LINEA DECORATIVA', 'PINTURA LATEX', 'AMERICAN COLORS', 'LATEX AMERICAN COLORS', 'BLANCO', '4GL'],
    ['4470', 'HERRAMIENTAS Y ACCESORIOS', 'BROCHAS', 'PONY', 'BROCHA PONY', '', '1 PLG'],
    ['5168', 'HERRAMIENTAS Y ACCESORIOS', 'BANDEJAS', 'C&A', 'BANDEJA PLASTICA C&A 9 PLG', '', ''],
]

INSTRUCTIONS = [
    'Llena la hoja «Productos» (es la que se importa). Cada fila es una variante: un producto en una presentación y un color.',
    'Columnas obligatorias: CATEGORIA, SUB CATEGORIA, MARCA y PRODUCTO. CODIGO, COLOR y PRESENTACION son opcionales, '
    'pero un producto nuevo solo se crea si su fila trae CODIGO.',
    'Escribe los nombres igual que en la web: elígelos de las listas desplegables o cópialos de la hoja «Listas». '
    'El producto se reconoce por su nombre y su marca: una marca distinta o un nombre cortado crean un producto nuevo.',
    'CODIGO: el código del artículo de esa presentación y ese color. Se conservan los ceros a la izquierda (0025).',
    'COLOR: el nombre del color. Si tiene código de fábrica, va al final con # (por ejemplo AMARILLO LIMON #41).',
    'PRESENTACION: igual que en la web (1GL, 4GL, 1/4GL, 12 PLG…). Revisa que ninguna celda quede con #N/A.',
    'Si un producto tiene colores, todas sus filas deben llevar color.',
    'Al importar, los colores y códigos de cada producto del archivo se reemplazan por los del archivo: '
    'incluye todas sus variantes, no solo las nuevas. Los productos que no estén en el archivo no cambian.',
    'Guarda el archivo como .xlsx y súbelo en Admin > Productos > Importar Excel.',
]


def import_header(value):
    """Encabezado de una columna sin tildes ni espacios de mas: 'CATEGORÍA' -> 'Categoria'."""
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode('ascii')
    return ' '.join(text.split()).title()


def _style_header(ws, headers, widths=None):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for i, width in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = 'A2'


def build_import_template(lists, examples=None):
    """Libro del modelo. lists: {'categorias', 'subcategorias', 'marcas', 'presentaciones': [nombres],
    'productos': [(nombre, marca, categoria)]}; examples: filas de ejemplo con las columnas de IMPORT_COLUMNS."""
    wb = Workbook()

    # --- Productos: la hoja que se importa ---
    ws = wb.active
    ws.title = 'Productos'
    _style_header(ws, IMPORT_COLUMNS, WIDTHS)
    for row in range(2, TEMPLATE_ROWS + 2):
        ws.cell(row=row, column=1).number_format = '@'   # CODIGO como texto: conserva los ceros

    # --- Listas: los nombres que ya usa la web ---
    ls = wb.create_sheet('Listas')
    products = sorted(lists.get('productos', []))
    columns = [
        ('CATEGORIAS', sorted(set(lists.get('categorias', [])))),
        ('SUB CATEGORIAS', sorted(set(lists.get('subcategorias', [])))),
        ('MARCAS', sorted(set(lists.get('marcas', [])))),
        ('PRESENTACIONES', sorted(set(lists.get('presentaciones', [])))),
        ('PRODUCTOS', [p[0] for p in products]),
        ('MARCA DEL PRODUCTO', [p[1] for p in products]),
        ('CATEGORIA DEL PRODUCTO', [p[2] for p in products]),
    ]
    _style_header(ls, [title for title, _ in columns], [24, 30, 22, 18, 46, 22, 26])
    for col, (_, values) in enumerate(columns, start=1):
        for row, value in enumerate(values, start=2):
            ls.cell(row=row, column=col, value=value)

    # Listas desplegables en Productos: avisan si el nombre no esta en la web, pero dejan seguir (puede ser nuevo)
    checks = [  # (columna en Productos, columna en Listas, que se crearia)
        ('B', 'A', 'una categoría nueva'), ('C', 'B', 'una subcategoría nueva'), ('D', 'C', 'una marca nueva'),
        ('E', 'E', 'un producto nuevo (si la fila trae CODIGO)'), ('G', 'D', 'una presentación nueva'),
    ]
    sizes = {letter: len(values) for letter, (_, values) in zip('ABCDEFG', columns)}
    for target, source, creates in checks:
        if not sizes[source]:
            continue
        dv = DataValidation(type='list', formula1=f'=Listas!${source}$2:${source}${sizes[source] + 1}', allow_blank=True,
                            showErrorMessage=True, errorStyle='warning', errorTitle='No está en la web',
                            error=f'Ese nombre no está en la web: al importar se creará {creates}. '
                                  'Revisa que no sea el mismo escrito distinto.')
        dv.add(f'{target}2:{target}{TEMPLATE_ROWS + 1}')
        ws.add_data_validation(dv)

    # --- Instrucciones con un ejemplo ---
    ins = wb.create_sheet('Instrucciones', 1)
    ins.column_dimensions['A'].width = 4
    ins['B1'] = 'Modelo de importación de productos — Grupo Dewill'
    ins['B1'].font = Font(bold=True, size=14, color='0F2B5B')
    for n, text in enumerate(INSTRUCTIONS, start=1):
        ins.cell(row=n + 2, column=1, value=n).font = Font(bold=True, color='D92027')
        ins.cell(row=n + 2, column=2, value=text).alignment = Alignment(wrap_text=True, vertical='top')
        ins.merge_cells(start_row=n + 2, start_column=2, end_row=n + 2, end_column=8)
        ins.row_dimensions[n + 2].height = 32
    start = len(INSTRUCTIONS) + 4
    ins.cell(row=start, column=2, value='Ejemplo (no se importa: copia el formato en la hoja «Productos»)').font = Font(bold=True)
    for i, title in enumerate(IMPORT_COLUMNS, start=2):
        cell = ins.cell(row=start + 1, column=i, value=title)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for r, row in enumerate(examples or DEFAULT_EXAMPLES, start=start + 2):
        for c, value in enumerate(row, start=2):
            ins.cell(row=r, column=c, value=value or None).number_format = '@'
    for i, width in enumerate(WIDTHS, start=2):
        ins.column_dimensions[ins.cell(row=1, column=i).column_letter].width = width

    wb.active = 0   # se abre en la hoja Productos
    return wb
